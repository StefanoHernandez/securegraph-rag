"""Generate deterministic MRD corpus PDF, DOCX, and XLSX artifacts."""

from __future__ import annotations

import argparse
import hashlib
import re
import zipfile
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path

from securgraph.corpus.manifest import load_corpus_manifest

FIXED_TIMESTAMP = (2026, 1, 1, 0, 0, 0)
FIXED_DATETIME = datetime(2026, 1, 1, tzinfo=UTC)
XLSX_DOC_IDS = frozenset({"10", "11", "12", "13"})
PDF_DOC_IDS = frozenset({"04", "05", "07", "17"})
GENERATED_DIR = "generated"


@dataclass(frozen=True)
class GeneratedArtifact:
    source: Path
    path: Path
    suffix: str


def expected_artifacts(corpus_root: Path) -> list[GeneratedArtifact]:
    """Return the stable generated artifact path for each manifest source document."""
    manifest = load_corpus_manifest(corpus_root)
    artifacts: list[GeneratedArtifact] = []
    for document in manifest.documents:
        doc_id = document.path.name.split("_", 1)[0]
        if doc_id in XLSX_DOC_IDS:
            suffix = ".xlsx"
        elif doc_id in PDF_DOC_IDS:
            suffix = ".pdf"
        else:
            suffix = ".docx"
        artifacts.append(
            GeneratedArtifact(
                source=document.path,
                path=manifest.root / GENERATED_DIR / f"{document.path.stem}{suffix}",
                suffix=suffix,
            )
        )
    return artifacts


def check_generated_artifacts(corpus_root: Path) -> list[str]:
    """Check that all deterministic artifacts exist and are non-empty."""
    errors: list[str] = []
    for artifact in expected_artifacts(corpus_root):
        if not artifact.path.is_file():
            errors.append(f"missing generated artifact: {artifact.path.relative_to(corpus_root)}")
        elif artifact.path.stat().st_size == 0:
            errors.append(f"empty generated artifact: {artifact.path.relative_to(corpus_root)}")
    return errors


def generate(corpus_root: Path) -> list[Path]:
    """Generate all MRD corpus artifacts from readable markdown sources."""
    outputs: list[Path] = []
    generated_dir = corpus_root / GENERATED_DIR
    generated_dir.mkdir(parents=True, exist_ok=True)
    for artifact in expected_artifacts(corpus_root):
        text = artifact.source.read_text(encoding="utf-8")
        if artifact.suffix == ".pdf":
            _write_pdf(artifact.path, artifact.source.stem, text)
        elif artifact.suffix == ".docx":
            _write_docx(artifact.path, artifact.source.stem, text)
        elif artifact.suffix == ".xlsx":
            _write_xlsx(artifact.path, artifact.source.stem, text)
        else:  # pragma: no cover - defensive guard for future mappings.
            raise ValueError(f"Unsupported generated artifact suffix: {artifact.suffix}")
        outputs.append(artifact.path)
    return outputs


def _markdown_title_and_body(text: str) -> tuple[str, list[str]]:
    lines = [line.rstrip() for line in text.splitlines()]
    title = "MRD Corpus Document"
    body: list[str] = []
    for line in lines:
        if line.startswith("# "):
            title = line.removeprefix("# ").strip() or title
        elif line.strip():
            body.append(line.strip())
    return title, body


def _write_pdf(path: Path, stem: str, text: str) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas

    title, body = _markdown_title_and_body(text)
    c = canvas.Canvas(str(path), pagesize=A4, invariant=1)
    c.setTitle(title)
    c.setAuthor("SecureGraph MRD corpus generator")
    _, height = A4
    x = 20 * mm
    y = height - 24 * mm
    c.setFont("Helvetica-Bold", 14)
    c.drawString(x, y, title)
    y -= 10 * mm
    c.setFont("Helvetica", 9)
    c.drawString(x, y, f"Source: {stem}.md")
    y -= 8 * mm
    c.setFont("Helvetica", 11)
    for paragraph in body:
        for line in _wrap_text(paragraph, 92):
            if y < 22 * mm:
                c.showPage()
                y = height - 24 * mm
                c.setFont("Helvetica", 11)
            c.drawString(x, y, line)
            y -= 6 * mm
        y -= 2 * mm
    c.save()


def _write_docx(path: Path, stem: str, text: str) -> None:
    import docx

    title, body = _markdown_title_and_body(text)
    document = docx.Document()
    document.core_properties.author = "SecureGraph MRD corpus generator"
    document.core_properties.created = FIXED_DATETIME
    document.core_properties.modified = FIXED_DATETIME
    document.core_properties.title = title
    document.add_heading(title, level=1)
    document.add_paragraph(f"Source: {stem}.md")
    for paragraph in body:
        document.add_paragraph(paragraph)
    document.save(path)
    _normalize_zip(path)


def _write_xlsx(path: Path, stem: str, text: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    title, body = _markdown_title_and_body(text)
    workbook = Workbook()
    workbook.properties.creator = "SecureGraph MRD corpus generator"
    workbook.properties.created = FIXED_DATETIME
    workbook.properties.modified = FIXED_DATETIME
    workbook.properties.title = title
    sheet = workbook.active
    sheet.title = "Source"
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    rows = [
        ("Field", "Value"),
        ("Title", title),
        ("Source", f"{stem}.md"),
        ("Body SHA256", hashlib.sha256(text.encode("utf-8")).hexdigest()),
    ]
    for row_index, row in enumerate(rows, start=1):
        sheet.append(row)
        if row_index == 1:
            for cell in sheet[row_index]:
                cell.fill = header_fill
                cell.font = header_font
    sheet.append(())
    sheet.append(("Line", "Text"))
    for cell in sheet[sheet.max_row]:
        cell.fill = header_fill
        cell.font = header_font
    for index, line in enumerate(body, start=1):
        sheet.append((index, line))
    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 110
    workbook.save(path)
    _normalize_zip(path)


def _wrap_text(text: str, width: int) -> list[str]:
    words = re.split(r"\s+", text.strip())
    lines: list[str] = []
    current = ""
    for word in words:
        candidate = f"{current} {word}".strip()
        if len(candidate) > width and current:
            lines.append(current)
            current = word
        else:
            current = candidate
    if current:
        lines.append(current)
    return lines


def _normalize_zip(path: Path) -> None:
    tmp = path.with_suffix(path.suffix + ".tmp")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
        tmp, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
    ) as target:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(filename=name, date_time=FIXED_TIMESTAMP)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            target.writestr(info, source.read(name))
    tmp.replace(path)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--corpus-root",
        type=Path,
        default=Path(__file__).resolve().parents[1] / "corpus" / "mrd_v1",
    )
    parser.add_argument("--check", action="store_true")
    args = parser.parse_args()
    if args.check:
        errors = check_generated_artifacts(args.corpus_root)
        for error in errors:
            print(error)
        return 1 if errors else 0
    outputs = generate(args.corpus_root)
    for output in outputs:
        print(output.relative_to(args.corpus_root))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
